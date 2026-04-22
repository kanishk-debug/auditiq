# core/invoice_processing/invoice_excel.py
# New file — not in Colab notebook
# Takes list of audit report dicts → writes color-coded Excel

import os
import tempfile
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side


# ── Color map by ITC status ────────────────────────────────────
STATUS_STYLE = {
    "ELIGIBLE":      ("C6EFCE", "276221"),   # green
    "BLOCKED":       ("FFDCDC", "9C0006"),   # red
    "REVIEW_NEEDED": ("FFEB9C", "9C6500"),   # yellow
    "RCM":           ("FFE0C0", "8B4000"),   # orange
}

HEADERS = [
    ("Vendor Name",        "vendor_name",        28),
    ("Vendor GSTIN",       "vendor_gstin",        18),
    ("Invoice No",         "invoice_number",      20),
    ("Invoice Date",       "invoice_date",        13),
    ("Taxable Value",      "taxable_value",       14),
    ("CGST",               "cgst",                10),
    ("SGST",               "sgst",                10),
    ("IGST",               "igst",                10),
    ("Total Amount",       "total_amount",        14),
    ("Type",               "invoice_type",         7),
    ("RCM",                "is_rcm",               6),
    ("ITC Status",         "itc_status",          14),
    ("ITC Amount",         "itc_amount",          12),
    ("ITC Reason",         "itc_reason",          45),
    ("Confidence",         "itc_confidence",      11),
    ("Extraction Score",   "extraction_confidence", 14),   # ← ADD THIS LINE
    ("Validation",         "validation_status",   13),
    ("Flags",              "flags",               28),
    ("Invoice ID",         "invoice_id",          36),
    ("File",               "filename",            22),
]


def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _font(hex_color, bold=False, size=9):
    return Font(color=hex_color, bold=bold, size=size, name="Calibri")

def _border():
    s = Side(style="thin", color="D0D0D0")
    return Border(left=s, right=s, top=s, bottom=s)

def _money(val):
    if val is None or str(val).lower() in ("none", ""):
        return "—"
    try:
        return f"Rs {float(val):,.2f}"
    except Exception:
        return str(val)


def generate_invoice_excel(audit_reports: list, output_path: str) -> str:
    """
    Takes a list of dicts from generate_audit_report().
    Writes a color-coded Excel to output_path.
    Returns output_path.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice Audit"
    ws.sheet_view.showGridLines = False

    # ── Header row ──
    for col_idx, (header, _, width) in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill      = _fill("2B2D42")
        cell.font      = _font("FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = _border()
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 24

    # ── Data rows ──
    for row_idx, report in enumerate(audit_reports, 2):
        status = report.get("itc_status", "REVIEW_NEEDED")
        fill_hex, font_hex = STATUS_STYLE.get(status, ("FFFFFF", "000000"))

        is_alt   = (row_idx % 2 == 0)
        row_fill = fill_hex if not is_alt else fill_hex  # same color, no alt needed for status rows

        for col_idx, (_, field, _) in enumerate(HEADERS, 1):
            raw = report.get(field, "")

            if field in ("taxable_value", "cgst", "sgst", "igst", "total_amount", "itc_amount", "gst_paid"):
                val = _money(raw)
            elif field == "is_rcm":
                val = "Yes" if raw else "No"
            elif field == "flags":
                val = str(raw) if raw else "None"
            else:
                val = str(raw) if raw is not None else ""

            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill      = _fill(fill_hex)
            cell.font      = _font(font_hex, size=9)
            cell.alignment = Alignment(
                vertical="center",
                horizontal="left",
                wrap_text=(field in ("itc_reason", "flags", "vendor_name"))
            )
            cell.border = _border()

        ws.row_dimensions[row_idx].height = 18

    ws.freeze_panes = "A2"

    # ── Summary block below data ──
    from collections import Counter
    statuses = Counter(r.get("itc_status", "REVIEW_NEEDED") for r in audit_reports)

    gap_row  = len(audit_reports) + 3
    ws.cell(row=gap_row, column=1, value="SUMMARY").font = _font("000000", bold=True, size=10)

    for i, (status, count) in enumerate(statuses.items()):
        r = gap_row + 1 + i
        fill_hex, font_hex = STATUS_STYLE.get(status, ("F1EFE8", "000000"))
        ws.cell(row=r, column=1, value=status).fill  = _fill(fill_hex)
        ws.cell(row=r, column=1).font                = _font(font_hex, bold=True)
        ws.cell(row=r, column=2, value=count).fill   = _fill(fill_hex)
        ws.cell(row=r, column=2).font                = _font(font_hex)

    wb.save(output_path)
    return output_path