# ============================================================
# AUDITIQ — PHASE 5E: RECONCILIATION EXCEL OUTPUT
# Produces a professional color-coded Excel for the CA
#
# Tabs:
#   1. Summary          — white   — ITC position + bucket totals
#   2. Matched          — green   — Claim ITC now
#   3. Ineligible       — red     — Do not claim (Sec 17(5))
#   4. Missing in 2B    — yellow  — Supplier hasn't filed
#   5. Missing in Books — blue    — Chase client
#   6. Previous Month   — purple  — Already handled
#   7. RCM              — orange  — Reverse charge entries
#   8. Review           — amber   — Needs manual CA check
#   9. Adjustments      — gray    — [NEW] Consolidated/bank charge entries  ← FIX S1
# ============================================================

from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.styles.numbers import FORMAT_DATE_DDMMYY
import pandas as pd
from datetime import datetime

# ── Tab config ─────────────────────────────────────────────────
TAB_CONFIG = {
    "Matched": {
        "buckets":     ["MATCHED", "PROBABLE_MATCH"],
        "header_fill": "1E7B34",
        "header_font": "FFFFFF",
        "alt_fill":    "EAF3DE",
        "tab_color":   "1E7B34",
        "action":      "Claim ITC this month",
    },
    "Ineligible": {
        "buckets":     ["INELIGIBLE"],
        "header_fill": "C00000",
        "header_font": "FFFFFF",
        "alt_fill":    "FFDCDC",
        "tab_color":   "C00000",
        "action":      "Do NOT claim ITC — blocked under Sec 17(5)",
    },
    "Missing in 2B": {
        "buckets":     ["MISSING_IN_2B"],
        "header_fill": "9C6500",
        "header_font": "FFFFFF",
        "alt_fill":    "FFEB9C",
        "tab_color":   "F0A500",
        "action":      "Supplier hasn't filed GSTR-1 — hold for next month",
    },
    "Missing in Books": {
        "buckets":     ["MISSING_IN_BOOKS"],
        "header_fill": "185FA5",
        "header_font": "FFFFFF",
        "alt_fill":    "DCEEFF",
        "tab_color":   "185FA5",
        "action":      "Invoice not in Tally — chase client or check",
    },
    "Previous Month": {
        "buckets":     ["PREVIOUS_MONTH"],
        "header_fill": "4B2580",
        "header_font": "FFFFFF",
        "alt_fill":    "EAE0F5",
        "tab_color":   "7B5EA7",
        "action":      "Already booked in prior month — no action",
    },
    "RCM": {
        "buckets":     ["RCM"],
        "header_fill": "8B4000",
        "header_font": "FFFFFF",
        "alt_fill":    "FFE0C0",
        "tab_color":   "E07020",
        "action":      "Pay GST to govt directly, then claim ITC",
    },
    "Review": {
        "buckets":     ["REVIEW"],
        "header_fill": "7B6800",
        "header_font": "FFFFFF",
        "alt_fill":    "FFF2CC",
        "tab_color":   "FFCC00",
        # ── FIX S3: updated action text to reference matched 2B invoice column ──
        "action":      "Invoice numbers differ — check 'Matched 2B Invoice' column, verify with vendor",
    },
    # ── FIX S1: ADJUSTMENT tab added — was missing, causing Bank Charges to be silently dropped ──
    "Adjustments": {
        "buckets":     ["ADJUSTMENT"],
        "header_fill": "3C3C3C",   # dark gray
        "header_font": "FFFFFF",
        "alt_fill":    "F1EFE8",   # light gray
        "tab_color":   "888780",
        "action":      "Consolidated / bank charge entries — verify ITC is claimed via journal entry",
    },
}

# ── Columns shown in every data tab ───────────────────────────
DATA_COLUMNS = [
    ("Vendor Name",          "vendor_name",         30),
    ("Invoice No",           "invoice_number",       22),
    ("Invoice Date",         "invoice_date",         13),
    ("Total Amount",         "total_amount",         14),
    ("IGST",                 "igst",                 12),
    ("CGST",                 "cgst",                 12),
    ("SGST",                 "sgst",                 12),
    ("Source",               "source",                7),
    ("Match Type",           "match_type",           16),
    ("Confidence",           "confidence",           10),
    # ── FIX S3: Matched 2B Invoice column added — was only buried in CA Note text ──
    ("Matched 2B Invoice",   "matched_invoice_2b",   22),
    ("CA Note",              "match_reasoning",      52),
]

# ── Style helpers ──────────────────────────────────────────────
def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color.lstrip("#"))

def _font(hex_color: str = "000000", bold: bool = False, size: int = 10) -> Font:
    return Font(color=hex_color.lstrip("#"), bold=bold, size=size,
                name="Calibri")

def _border(color: str = "D0D0D0") -> Border:
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def _center(wrap: bool = False) -> Alignment:
    return Alignment(horizontal="center", vertical="center",
                     wrap_text=wrap)

def _left(wrap: bool = False) -> Alignment:
    return Alignment(horizontal="left", vertical="center",
                     wrap_text=wrap)

def _money(val) -> str:
    try:
        v = float(val or 0)
        if v == 0:
            return "—"
        return f"₹{v:,.2f}"
    except Exception:
        return "—"

def _date_str(val) -> str:
    try:
        if pd.isna(val):
            return ""
        return pd.Timestamp(val).strftime("%d/%m/%Y")
    except Exception:
        return str(val) if val else ""

# ── FIX S3: helper to detect credit notes / reversals ─────────
def _is_reversal(row) -> bool:
    """Returns True if any tax column is negative — indicates a credit note or reversal."""
    try:
        igst = float(row.get("igst", 0) or 0)
        cgst = float(row.get("cgst", 0) or 0)
        sgst = float(row.get("sgst", 0) or 0)
        inv_type = str(row.get("invoice_type", "") or "").strip().lower()
        return igst < 0 or cgst < 0 or sgst < 0 or "credit" in inv_type
    except Exception:
        return False


# ── Summary tab ────────────────────────────────────────────────
def _write_summary(ws, df: pd.DataFrame, period: str):
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:G1")
    ws["A1"].value     = "AuditIQ — GST Input Tax Credit Reconciliation"
    ws["A1"].font      = _font("1A1A2E", bold=True, size=16)
    ws["A1"].alignment = _center()
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:G2")
    ws["A2"].value     = f"Period: {period}   |   Generated: {datetime.now().strftime('%d %b %Y, %I:%M %p')}"
    ws["A2"].font      = _font("666666", size=10)
    ws["A2"].alignment = _center()
    ws.row_dimensions[2].height = 18

    row = 4
    ws.merge_cells(f"A{row}:G{row}")
    ws[f"A{row}"].value     = "ITC POSITION"
    ws[f"A{row}"].font      = _font("1A1A2E", bold=True, size=11)
    ws[f"A{row}"].fill      = _fill("F0F0F0")
    ws[f"A{row}"].alignment = _left()
    ws.row_dimensions[row].height = 22

    def itc_row(label, mask, fill_hex, font_hex, row_num):
        sub = df[mask]
        igst  = round(sub["igst"].fillna(0).astype(float).sum(), 2)
        cgst  = round(sub["cgst"].fillna(0).astype(float).sum(), 2)
        sgst  = round(sub["sgst"].fillna(0).astype(float).sum(), 2)
        total = round(igst + cgst + sgst, 2)
        vals  = [label, len(sub), _money(igst), _money(cgst), _money(sgst), _money(total), ""]
        for col_idx, val in enumerate(vals, 1):
            c = ws.cell(row=row_num, column=col_idx, value=val)
            c.fill      = _fill(fill_hex)
            c.font      = _font(font_hex, bold=(col_idx == 1), size=10)
            c.alignment = _center() if col_idx > 1 else _left()
            c.border    = _border()
        ws.row_dimensions[row_num].height = 20

    row = 5
    hdr = ["Category", "Count", "IGST", "CGST", "SGST", "Total Tax", ""]
    for ci, h in enumerate(hdr, 1):
        c = ws.cell(row=row, column=ci, value=h)
        c.fill = _fill("2B2D42"); c.font = _font("FFFFFF", bold=True, size=9)
        c.alignment = _center(); c.border = _border()
    ws.row_dimensions[row].height = 18

    row = 6
    itc_row("✅ Claimable ITC",
            df["bucket"].isin(["MATCHED", "PROBABLE_MATCH"]),
            "C6EFCE", "276221", row); row += 1
    itc_row("🚫 Blocked ITC",
            df["bucket"] == "INELIGIBLE",
            "FFDCDC", "9C0006", row); row += 1
    itc_row("🔄 RCM Payable",
            df["bucket"] == "RCM",
            "FFE0C0", "8B4000", row); row += 1
    itc_row("⏳ On Hold (Missing in 2B)",
            df["bucket"] == "MISSING_IN_2B",
            "FFEB9C", "9C6500", row); row += 1
    itc_row("📋 Pending (Missing in Books)",
            df["bucket"] == "MISSING_IN_BOOKS",
            "DCEEFF", "185FA5", row); row += 1
    itc_row("📅 Previous Month",
            df["bucket"] == "PREVIOUS_MONTH",
            "EAE0F5", "4B2580", row); row += 1
    itc_row("⚠️ Needs Review",
            df["bucket"].isin(["REVIEW", "ADJUSTMENT"]),
            "FFF2CC", "9C6500", row); row += 1

    all_igst = round(df["igst"].fillna(0).astype(float).sum(), 2)
    all_cgst = round(df["cgst"].fillna(0).astype(float).sum(), 2)
    all_sgst = round(df["sgst"].fillna(0).astype(float).sum(), 2)
    total_vals = ["TOTAL", len(df), _money(all_igst), _money(all_cgst),
                  _money(all_sgst), _money(all_igst + all_cgst + all_sgst), ""]
    for ci, val in enumerate(total_vals, 1):
        c = ws.cell(row=row, column=ci, value=val)
        c.fill = _fill("E8E8E8"); c.font = _font("000000", bold=True, size=10)
        c.alignment = _center() if ci > 1 else _left()
        c.border = _border()
    ws.row_dimensions[row].height = 20
    row += 2

    ws.merge_cells(f"A{row}:G{row}")
    ws[f"A{row}"].value     = "ITEMS REQUIRING CA ACTION"
    ws[f"A{row}"].font      = _font("1A1A2E", bold=True, size=11)
    ws[f"A{row}"].fill      = _fill("F0F0F0")
    ws[f"A{row}"].alignment = _left()
    ws.row_dimensions[row].height = 22
    row += 1

    action_buckets = ["MISSING_IN_BOOKS", "MISSING_IN_2B", "REVIEW"]
    action_df = df[df["bucket"].isin(action_buckets)].copy()

    if not action_df.empty:
        for ci, h in enumerate(["Vendor", "Invoice No", "Bucket", "Tax Amount", "Action", "", ""], 1):
            c = ws.cell(row=row, column=ci, value=h)
            c.fill = _fill("2B2D42"); c.font = _font("FFFFFF", bold=True, size=9)
            c.alignment = _center(); c.border = _border()
        ws.row_dimensions[row].height = 18
        row += 1

        bucket_fills = {
            "MISSING_IN_BOOKS": ("DCEEFF", "185FA5"),
            "MISSING_IN_2B":    ("FFEB9C", "9C6500"),
            "REVIEW":           ("FFF2CC", "9C6500"),
        }
        for _, arow in action_df.iterrows():
            igst = float(arow.get("igst", 0) or 0)
            cgst = float(arow.get("cgst", 0) or 0)
            tax  = _money(igst if igst else cgst)
            fill_h, font_h = bucket_fills.get(arow.get("bucket", ""), ("F5F5F5", "000000"))

            # ── FIX S3: credit note gets its own action label in summary action table ──
            if arow.get("bucket") == "MISSING_IN_BOOKS" and _is_reversal(arow):
                action_label = "⚠️ Credit note / reversal — verify if booked in Tally"
            else:
                action_label = {
                    "MISSING_IN_BOOKS": "Chase client / check Tally",
                    "MISSING_IN_2B":    "Wait — supplier not filed",
                    "REVIEW":           "Verify invoice number — see Matched 2B Invoice column",
                }.get(arow.get("bucket", ""), "")

            vals = [
                str(arow.get("vendor_name", ""))[:42],
                str(arow.get("invoice_number", ""))[:22],
                arow.get("bucket", ""),
                tax,
                action_label, "", ""
            ]
            for ci, val in enumerate(vals, 1):
                c = ws.cell(row=row, column=ci, value=val)
                c.fill = _fill(fill_h); c.font = _font(font_h, size=9)
                c.alignment = _left(); c.border = _border()
            ws.row_dimensions[row].height = 17
            row += 1

    for col, width in zip(["A", "B", "C", "D", "E", "F", "G"],
                          [36, 8, 14, 14, 14, 16, 4]):
        ws.column_dimensions[col].width = width


# ── Data tab writer ────────────────────────────────────────────
def _write_data_tab(ws, df: pd.DataFrame, config: dict, tab_name: str):
    ws.sheet_tab_color = config.get("tab_color", "FFFFFF")
    ws.sheet_view.showGridLines = False

    header_fill = config["header_fill"]
    header_font = config["header_font"]
    alt_fill    = config["alt_fill"]

    col_count = len(DATA_COLUMNS)
    ws.merge_cells(f"A1:{get_column_letter(col_count)}1")
    ws["A1"].value     = f"{tab_name}  —  {config['action']}  ({len(df)} entries)"
    ws["A1"].font      = _font(header_font, bold=True, size=11)
    ws["A1"].fill      = _fill(header_fill)
    ws["A1"].alignment = _left()
    ws.row_dimensions[1].height = 24

    for col_idx, (col_label, _, _width) in enumerate(DATA_COLUMNS, start=1):
        c = ws.cell(row=2, column=col_idx, value=col_label)
        c.fill      = _fill(header_fill)
        c.font      = _font(header_font, bold=True, size=9)
        c.alignment = _center(wrap=True)
        c.border    = _border(header_fill)
    ws.row_dimensions[2].height = 24

    if df.empty:
        ws.merge_cells(f"A3:{get_column_letter(col_count)}3")
        ws["A3"].value     = "No entries in this category for the period."
        ws["A3"].font      = _font("888888", size=10)
        ws["A3"].alignment = _center()
        ws.row_dimensions[3].height = 22
        return

    for row_idx, (_, row) in enumerate(df.iterrows(), start=3):
        is_alt = (row_idx % 2 == 1)

        # ── FIX S3: credit note / reversal rows get a distinct amber fill ──
        # so CA can immediately spot them without reading the CA Note text
        if _is_reversal(row):
            row_fill = _fill("FFD580")   # amber highlight — flags reversal visually
        else:
            row_fill = _fill(alt_fill) if is_alt else _fill("FFFFFF")

        for col_idx, (_, col_key, _) in enumerate(DATA_COLUMNS, start=1):
            raw = row.get(col_key, "")

            if col_key in ("total_amount", "igst", "cgst", "sgst"):
                val   = _money(raw)
                align = _center()

            elif col_key == "invoice_date":
                val   = _date_str(raw)
                align = _center()

            elif col_key == "match_reasoning":
                val   = str(raw)[:200] if pd.notna(raw) else ""
                align = _left(wrap=True)

            elif col_key == "matched_invoice_2b":
                # ── FIX S3: show the 2B invoice number consumed by fuzzy match ──
                # Previously this only appeared buried inside CA Note text
                val   = str(raw) if pd.notna(raw) and str(raw) not in ("", "nan", "None") else "—"
                align = _left()

            elif col_key == "source":
                val   = str(raw) if pd.notna(raw) else ""
                align = _center()

            elif col_key == "confidence":
                val   = str(raw) if pd.notna(raw) else ""
                align = _center()
                # Confidence cell gets its own color override (HIGH=green, LOW=red)
                # but credit note rows keep amber to avoid confusion — amber wins
                if _is_reversal(row):
                    conf_fill = row_fill
                elif val == "HIGH":
                    conf_fill = _fill("EAF3DE")
                elif val == "LOW":
                    conf_fill = _fill("FFDCDC")
                else:
                    conf_fill = row_fill
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.fill = conf_fill
                c.font = _font("000000", size=9)
                c.alignment = align
                c.border = _border()
                continue

            else:
                val   = str(raw) if pd.notna(raw) else ""
                align = _left()

            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.fill      = row_fill
            c.font      = _font("000000", size=9)
            c.alignment = align
            c.border    = _border()

        ws.row_dimensions[row_idx].height = 18 if col_key != "match_reasoning" else 32

    for col_idx, (_, _, width) in enumerate(DATA_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A3"

    # Tax totals footer
    footer_row = len(df) + 4
    igst_sum = round(df["igst"].fillna(0).astype(float).sum(), 2)
    cgst_sum = round(df["cgst"].fillna(0).astype(float).sum(), 2)
    sgst_sum = round(df["sgst"].fillna(0).astype(float).sum(), 2)

    footer_vals = {1: "TOTAL", 5: _money(igst_sum), 6: _money(cgst_sum), 7: _money(sgst_sum)}
    for ci in range(1, col_count + 1):
        c = ws.cell(row=footer_row, column=ci, value=footer_vals.get(ci, ""))
        c.fill      = _fill(header_fill)
        c.font      = _font(header_font, bold=True, size=9)
        c.alignment = _center()
        c.border    = _border(header_fill)
    ws.row_dimensions[footer_row].height = 20


# ── Entry Point ────────────────────────────────────────────────
def generate_reconciliation_excel(
    recon_result: dict,
    output_path: str,
    period: str = ""
) -> str:
    """
    Takes reconcile() output.
    Writes a color-coded multi-tab Excel file.
    Returns the output path.

    Usage:
        output_file = generate_reconciliation_excel(
            recon_result = recon,
            output_path  = "/content/AuditIQ_Reconciliation_Feb2026.xlsx",
            period       = "February 2026"
        )
    """
    df = recon_result["data"].copy()
    wb = Workbook()

    # ── Tab 1: Summary ──
    _write_summary(wb.active, df, period)

    # ── Tabs 2-9: Data tabs ──
    tabs_written = 0
    total_rows   = 0

    for tab_name, config in TAB_CONFIG.items():
        subset = df[df["bucket"].isin(config["buckets"])].copy()
        ws     = wb.create_sheet(title=tab_name)
        _write_data_tab(ws, subset, config, tab_name)
        tabs_written += 1
        total_rows   += len(subset)

    wb.save(output_path)

    print(f"\n{'='*60}")
    print(f"✅ RECONCILIATION EXCEL GENERATED")
    print(f"{'='*60}")
    print(f"  File    : {output_path}")
    print(f"  Period  : {period}")
    print(f"  Tabs    : Summary + {tabs_written} data tabs")
    print(f"  Rows    : {total_rows}")
    print(f"  Size    : {__import__('os').path.getsize(output_path)/1024:.1f} KB")
    print(f"{'='*60}")
    print(f"  📁 Download from Colab Files panel (left sidebar)")
    print(f"{'='*60}\n")

    return output_path


